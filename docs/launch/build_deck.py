"""Build the marshal.build launch deck (.pptx) from the Marp source next to this file.

Source of truth is docs/launch/launch-deck.md: slides are separated by `---`, HTML
comments are speaker notes (Marp directive comments such as `<!-- _class: ... -->` are
skipped), markdown tables become real pptx tables, and the two diagram slides
("one picture" and "Architecture") are drawn as box-and-arrow shapes. Any other fenced
code block falls back to a monospace text panel, so nothing renders as a broken image.

Usage (scratch venv, nothing is installed into the repo):
  $HOME/.local/bin/uv venv /tmp/pptx-venv
  $HOME/.local/bin/uv pip install --python /tmp/pptx-venv/bin/python python-pptx
  /tmp/pptx-venv/bin/python docs/launch/build_deck.py
"""

import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "launch-deck.md"
OUTPUT = HERE / "marshal-build-launch-deck.pptx"

# Palette (Tailwind names for reference).
BG = RGBColor(0x0F, 0x17, 0x2A)  # slate-900
BODY = RGBColor(0xE2, 0xE8, 0xF0)  # slate-200
TITLE = RGBColor(0xF8, 0xFA, 0xFC)  # slate-50
ACCENT = RGBColor(0x81, 0x8C, 0xF8)  # indigo-400
MUTED = RGBColor(0x94, 0xA3, 0xB8)  # slate-400
PANEL = RGBColor(0x1E, 0x29, 0x3B)  # slate-800
PANEL_ACCENT = RGBColor(0x31, 0x2E, 0x81)  # indigo-900
CODE = RGBColor(0xC7, 0xD2, 0xFE)  # indigo-200

TITLE_FONT = "Calibri"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
SYMBOL_FONT = "Segoe UI Symbol"  # ships with Office; Calibri has no ★ glyph

SLIDE_W = Emu(12192000)  # 13.333 in (16:9)
SLIDE_H = Emu(6858000)  # 7.5 in
SLIDE_W_IN = 13.333
MARGIN = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN
FOOTER_TOP = Inches(6.95)

# ----------------------------------------------------------------------------- parsing

Block = tuple[str, object]  # ("para", str) | ("bullets", list[(level, text)]) | ("table", rows) ...


@dataclass
class SlideData:
    title: str = ""
    level: int = 2  # 1 = cover (h1), 2 = normal (h2)
    subtitle: str = ""
    blocks: list[Block] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def blocks_of(self, kind: str) -> list:
        return [payload for k, payload in self.blocks if k == kind]


DIRECTIVE = re.compile(
    r"^<!--\s*_?(class|paginate|header|footer|backgroundColor|backgroundImage|color)\s*:.*-->$"
)
COMMENT = re.compile(r"^<!--\s*(.*?)\s*-->$", re.S)
IMAGE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
TABLE_RULE = re.compile(r":?-+:?")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*|★)")
ESCAPES = re.compile(r"\\([\\`*_{}\[\]()#+\-.!<>|])")


def parse_deck(text: str) -> list[SlideData]:
    body = text
    if body.startswith("---\n"):
        end = body.index("\n---\n", 4)
        body = body[end + len("\n---\n") :]
    return [parse_slide(chunk) for chunk in re.split(r"\n---\n", body) if chunk.strip()]


def _is_special(line: str) -> bool:
    s = line.strip()
    return not s or s.startswith(("#", "- ", "|", "```", "<!--", "!["))


def parse_slide(chunk: str) -> SlideData:
    slide = SlideData()
    lines = chunk.strip("\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("<!--"):
            m = COMMENT.match(s)
            if m and not DIRECTIVE.match(s):
                slide.notes.append(m.group(1))
            i += 1
            continue
        if s.startswith("```"):
            j = i + 1
            code: list[str] = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                code.append(lines[j])
                j += 1
            slide.blocks.append(("code", "\n".join(code)))
            i = j + 1
            continue
        if s.startswith("# "):
            slide.title, slide.level = s[2:].strip(), 1
        elif s.startswith("## "):
            slide.title, slide.level = s[3:].strip(), 2
        elif s.startswith("### "):
            slide.subtitle = s[4:].strip()
        elif s.startswith("!["):
            m = IMAGE.match(s)
            if m:
                slide.blocks.append(("image", m.group(1)))
        elif s.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(TABLE_RULE.fullmatch(c) for c in cells):
                    rows.append(cells)
                i += 1
            slide.blocks.append(("table", rows))
            continue
        elif s.startswith("- "):
            items: list[tuple[int, str]] = []
            while i < len(lines) and lines[i].strip().startswith("- "):
                indent = len(lines[i]) - len(lines[i].lstrip())
                items.append((1 if indent >= 2 else 0, lines[i].strip()[2:].strip()))
                i += 1
            slide.blocks.append(("bullets", items))
            continue
        else:
            para = [s]
            while i + 1 < len(lines) and not _is_special(lines[i + 1]):
                i += 1
                para.append(lines[i].strip())
            slide.blocks.append(("para", " ".join(para)))
        i += 1
    return slide


def unescape(text: str) -> str:
    return ESCAPES.sub(r"\1", text)


def inline_runs(text: str) -> list[tuple[str, dict]]:
    """Split markdown inline formatting into (text, style) runs."""
    runs: list[tuple[str, dict]] = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((unescape(part[2:-2]), {"bold": True}))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append((part[1:-1], {"code": True}))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append((unescape(part[1:-1]), {"italic": True}))
        elif part == "★":
            runs.append((part, {"marker": True}))  # owner re-checks these on release day
        else:
            runs.append((unescape(part), {}))
    return runs


def plain(text: str) -> str:
    return "".join(t for t, _ in inline_runs(text))


# ----------------------------------------------------------------------------- primitives


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    return slide


def add_textbox(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    return box, tf


def fill_runs(paragraph, runs, size, *, color=BODY, font=BODY_FONT, bold=False, italic=False):
    for text, style in runs:
        run = paragraph.add_run()
        run.text = text
        f = run.font
        f.bold = bold or style.get("bold", False)
        f.italic = italic or style.get("italic", False)
        if style.get("code"):
            f.name = MONO_FONT
            f.color.rgb = CODE
            f.size = Pt(max(size - 3, 12))  # Consolas sits visually larger than Calibri
        elif style.get("marker"):
            f.name = SYMBOL_FONT
            f.color.rgb = ACCENT
            f.size = Pt(size)
        else:
            f.name = font
            f.color.rgb = color
            f.size = Pt(size)


def set_bullet(paragraph, level: int) -> None:
    """python-pptx has no bullet API; write the a:pPr children directly (schema order)."""
    pPr = paragraph._p.get_or_add_pPr()
    step = Inches(0.32)
    pPr.set("marL", str(int(step * (level + 1))))
    pPr.set("indent", str(-int(step)))
    buClr = pPr.makeelement(qn("a:buClr"), {})
    buClr.append(pPr.makeelement(qn("a:srgbClr"), {"val": str(ACCENT)}))
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"}))


def estimate_height_in(items: list[tuple[int, str, int]], size: int, width_in: float) -> float:
    """Rough stacked height (inches) for (level, text, space_after_pt) items at `size` pt."""
    total = 0.0
    for level, text, after in items:
        usable = width_in - 0.4 * (level + 1)
        per_line = max(1, int(usable * 72 / (size * 0.5)))
        lines = max(1, math.ceil(len(plain(text)) / per_line))
        total += lines * size * 1.2 / 72 + after / 72
    return total * 1.06 + 0.12


def add_title(slide, text: str):
    """Title at the top; returns the top edge (Emu) for body content."""
    long_title = len(text) > 48
    size = 32 if long_title else 36
    height = Inches(1.15) if long_title else Inches(0.8)
    box, tf = add_textbox(slide, MARGIN, Inches(0.35), CONTENT_W, height, anchor=MSO_ANCHOR.BOTTOM)
    box.name = "Title"
    p = tf.paragraphs[0]
    number, sep, rest = text.partition(" · ")
    if sep and number.isdigit():
        fill_runs(p, [(number, {})], size, color=ACCENT, font=TITLE_FONT, bold=True)
        fill_runs(p, inline_runs(sep + rest), size, color=TITLE, font=TITLE_FONT, bold=True)
    else:
        fill_runs(p, inline_runs(text), size, color=TITLE, font=TITLE_FONT, bold=True)
    rule_top = Inches(0.35) + height + Inches(0.1)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, rule_top, Inches(1.3), Inches(0.05))
    rule.fill.solid()
    rule.fill.fore_color.rgb = ACCENT
    rule.line.fill.background()
    rule.shadow.inherit = False
    return rule_top + Inches(0.3)


def add_footer(slide, index: int) -> None:
    box, tf = add_textbox(slide, MARGIN, FOOTER_TOP, Inches(4), Inches(0.3))
    fill_runs(tf.paragraphs[0], [("marshal.build", {})], 11, color=MUTED)
    box, tf = add_textbox(slide, SLIDE_W - MARGIN - Inches(1), FOOTER_TOP, Inches(1), Inches(0.3))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fill_runs(p, [(str(index), {})], 11, color=MUTED)


def draw_box(slide, x, y, w, h, text, *, size=14, fill=PANEL, line=ACCENT, color=BODY, bold=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.18
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    fill_runs(p, inline_runs(text), size, color=color, bold=bold)
    return shape


def draw_line(slide, x1, y1, x2, y2, *, start_arrow=False, end_arrow=False, color=ACCENT):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(1.5)
    ln = conn.line._get_or_add_ln()
    if start_arrow:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if end_arrow:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return conn


def draw_arrow(slide, x1, y1, x2, y2, *, both=False):
    return draw_line(slide, x1, y1, x2, y2, start_arrow=both, end_arrow=True)


def draw_label(slide, x, y, w, text, *, size=13):
    box, tf = add_textbox(slide, Inches(x), Inches(y), Inches(w), Inches(0.3))
    fill_runs(tf.paragraphs[0], [(text, {})], size, color=MUTED, italic=True)


# ----------------------------------------------------------------------------- content blocks


def render_text(slide, kind, payload, top, size) -> int:
    """One text box for a paragraph or a bullet list; returns the bottom edge (Emu)."""
    if kind == "para":
        items = [(0, payload, 10)]
    else:
        items = [(level, text, 8 if level == 0 else 5) for level, text in payload]
    height = Inches(estimate_height_in(items, size, CONTENT_W / 914400))
    box, tf = add_textbox(slide, MARGIN, top, CONTENT_W, height)
    for n, (level, text, after) in enumerate(items):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.space_after = Pt(after)
        p.line_spacing = 1.05
        if kind == "bullets":
            set_bullet(p, level)
        fill_runs(p, inline_runs(text), size - 2 * level)
    return top + height


def render_table(slide, rows: list[list[str]], top) -> int:
    n_rows, n_cols = len(rows), max(len(r) for r in rows)
    row_h = Inches(0.52)
    shape = slide.shapes.add_table(n_rows, n_cols, MARGIN, top, CONTENT_W, row_h * n_rows)
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    if n_cols == 2:
        table.columns[0].width = Inches(4.6)
        table.columns[1].width = CONTENT_W - Inches(4.6)
    for r, row in enumerate(rows):
        table.rows[r].height = row_h
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL_ACCENT if r == 0 else PANEL
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            text = row[c] if c < len(row) else ""
            fill_runs(
                cell.text_frame.paragraphs[0],
                inline_runs(text),
                14,
                color=TITLE if r == 0 else BODY,
                bold=(r == 0),
            )
    return top + row_h * n_rows


def render_code(slide, code: str, top) -> int:
    """Fallback for fenced blocks that have no dedicated diagram: a monospace panel."""
    lines = code.split("\n")
    size = 13
    height = Inches(len(lines) * size * 1.25 / 72 + 0.3)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, top, CONTENT_W, height)
    shape.adjustments[0] = 0.05
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.2)
    tf.margin_top = tf.margin_bottom = Inches(0.12)
    for n, line in enumerate(lines):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        fill_runs(p, [(line or " ", {"code": True})], size + 1)
    return top + height


def render_blocks(slide, blocks: list[Block], top, size: int) -> int:
    y = top
    for kind, payload in blocks:
        if kind == "table":
            y = render_table(slide, payload, y) + Inches(0.3)
        elif kind == "code":
            y = render_code(slide, payload, y) + Inches(0.2)
        elif kind in ("para", "bullets"):
            y = render_text(slide, kind, payload, y, size) + Inches(0.05)
    return y


# ----------------------------------------------------------------------------- diagrams


def draw_pipeline(slide, top) -> int:
    """Slide 3: Build (top) -> Govern (middle) -> Deploy (bottom), one spine of arrows."""
    y = top / 914400 + 0.05
    h, gap, spine_x = 0.6, 0.5, 1.5

    draw_box(slide, 0.6, y, 3.6, h, "Chat / Guided wizard / Studio")
    draw_arrow(slide, spine_x, y + h, spine_x, y + h + gap)
    draw_label(slide, spine_x + 0.25, y + h + 0.1, 6, "(PDF · DOCX · pasted specs as grounding)")

    y += h + gap
    draw_box(slide, 0.6, y, 5.0, h, "Spec set: requirements.md · design.md · tasks.md")
    draw_arrow(slide, 5.6, y + h / 2, 6.2, y + h / 2)
    draw_box(slide, 6.2, y, 3.9, h, "versions · diff · rollback · comments")
    draw_arrow(slide, spine_x, y + h, spine_x, y + h + gap)

    y += h + gap
    draw_box(slide, 0.6, y, 2.2, h, "Code generation")
    draw_arrow(slide, 2.8, y + h / 2, 3.3, y + h / 2)
    draw_box(slide, 3.3, y, 3.3, h, "deterministic validation gate")
    draw_arrow(slide, 6.6, y + h / 2, 7.1, y + h / 2)
    draw_box(slide, 7.1, y, 2.8, h, "spec-conformance review")
    draw_label(
        slide, 3.3, y + h + 0.1, 7, "(allowlist · packaging · auth contract · guardrail patterns)"
    )
    draw_arrow(slide, spine_x, y + h, spine_x, y + h + gap)

    y += h + gap
    draw_box(slide, 0.6, y, 3.3, h, "Risk scoring (7-factor rubric)")
    draw_arrow(slide, 3.9, y + h / 2, 4.4, y + h / 2)
    draw_box(slide, 4.4, y, 3.7, h, "auto-approve / maker-checker review")
    draw_arrow(slide, spine_x, y + h, spine_x, y + h + gap)

    y += h + gap
    draw_box(
        slide,
        0.6,
        y,
        8.0,
        h,
        "Enclave: a vended AWS account per deployment · TTL · budget · health · teardown",
        fill=PANEL_ACCENT,
        bold=True,
    )
    return Inches(y + h)


def draw_architecture(slide, top) -> int:
    """Slide 10: request path across the top, FastAPI fans out over a bus to its services."""
    y1 = top / 914400 + 0.05
    chain = ["Browser", "CloudFront + WAF", "Next.js", "FastAPI"]
    h1, w1, gap1 = 0.55, 2.3, 0.55
    x = (SLIDE_W_IN - (len(chain) * w1 + (len(chain) - 1) * gap1)) / 2
    api_x = 0.0
    for n, label in enumerate(chain):
        is_api = label == "FastAPI"
        draw_box(slide, x, y1, w1, h1, label, fill=PANEL_ACCENT if is_api else PANEL, bold=is_api)
        if is_api:
            api_x = x + w1 / 2
        if n < len(chain) - 1:
            draw_arrow(slide, x + w1, y1 + h1 / 2, x + w1 + gap1, y1 + h1 / 2)
        x += w1 + gap1

    # (label, bidirectional link from FastAPI, optional (downstream label, bidirectional)).
    spokes = [
        ("Postgres", True, None),
        ("DynamoDB (chat, shared state)", True, None),
        ("Bedrock seam (clamps · redaction · recording)", False, ("Amazon Bedrock", False)),
        (
            "Codegen (in-process, or the CodeBuild workspace runner)",
            False,
            ("S3 workspace/artifacts", True),
        ),
        ("Risk gate + review", False, None),
        (
            "Innovation Sandbox (Enclave vending)",
            False,
            ("Enclave account running the CloudFormation stack", False),
        ),
    ]
    n = len(spokes)
    side, gap2, h2, h3 = 0.5, 0.22, 0.85, 0.75
    w2 = (SLIDE_W_IN - 2 * side - (n - 1) * gap2) / n
    bus_y = y1 + h1 + 0.45
    y2 = bus_y + 0.45
    y3 = y2 + h2 + 0.45
    centers = [side + i * (w2 + gap2) + w2 / 2 for i in range(n)]

    draw_line(slide, api_x, y1 + h1, api_x, bus_y)
    draw_line(slide, min(centers[0], api_x), bus_y, max(centers[-1], api_x), bus_y)
    for cx, (label, both, downstream) in zip(centers, spokes, strict=True):
        draw_arrow(slide, cx, bus_y, cx, y2, both=both)
        draw_box(slide, cx - w2 / 2, y2, w2, h2, label, size=12)
        if downstream:
            text, both2 = downstream
            draw_arrow(slide, cx, y2 + h2, cx, y3, both=both2)
            is_enclave = text.startswith("Enclave")
            draw_box(
                slide,
                cx - w2 / 2,
                y3,
                w2,
                h3,
                text,
                size=12,
                fill=PANEL_ACCENT if is_enclave else PANEL,
                bold=is_enclave,
            )
    return Inches(y3 + h3)


# ----------------------------------------------------------------------------- slide types


def render_cover(slide, data: SlideData) -> None:
    y = Inches(0.7)
    for rel in data.blocks_of("image"):
        path = (SOURCE.parent / rel).resolve()
        if path.exists():
            size = Inches(1.7)
            pic = slide.shapes.add_picture(str(path), (SLIDE_W - size) // 2, y, size, size)
            pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE  # app-icon corners
            y += size + Inches(0.25)
            break

    box, tf = add_textbox(slide, MARGIN, y, CONTENT_W, Inches(0.85), anchor=MSO_ANCHOR.MIDDLE)
    box.name = "Title"
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    fill_runs(p, inline_runs(data.title), 40, color=TITLE, font=TITLE_FONT, bold=True)
    y += Inches(0.85)

    if data.subtitle:
        box, tf = add_textbox(slide, MARGIN, y, CONTENT_W, Inches(0.6), anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        fill_runs(p, inline_runs(data.subtitle), 26, color=ACCENT, font=TITLE_FONT, bold=True)
        y += Inches(0.8)

    width = Inches(10.6)
    for n, text in enumerate(data.blocks_of("para")):
        size, color = (18, BODY) if n == 0 else (16, MUTED)
        height = Inches(estimate_height_in([(0, text, 0)], size, width / 914400))
        box, tf = add_textbox(slide, (SLIDE_W - width) // 2, y, width, height)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.1
        fill_runs(p, inline_runs(text), size, color=color)
        y += height + Inches(0.2)


def render_hero(slide, data: SlideData, top) -> None:
    """A title plus one statement, centred in the remaining space (slide 14)."""
    height = FOOTER_TOP - top
    box, tf = add_textbox(slide, MARGIN, top, CONTENT_W, height, anchor=MSO_ANCHOR.MIDDLE)
    for n, text in enumerate(data.blocks_of("para")):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.15
        fill_runs(p, inline_runs(text), 28)


def body_size(data: SlideData) -> int:
    chars = sum(
        len(plain(t))
        for kind, payload in data.blocks
        if kind in ("para", "bullets")
        for t in ([payload] if kind == "para" else [text for _, text in payload])
    )
    return 20 if chars < 450 else 18


def render_slide(prs, data: SlideData, index: int) -> None:
    slide = new_slide(prs)
    if data.level == 1:
        render_cover(slide, data)
    else:
        top = add_title(slide, data.title)
        title_lc = data.title.lower()
        has_code = bool(data.blocks_of("code"))
        text_only = all(kind == "para" for kind, _ in data.blocks)
        if has_code and "one picture" in title_lc:
            draw_pipeline(slide, top)
        elif "architecture" in title_lc and data.blocks_of("para"):
            # The prose flow description is rendered as boxes and arrows; bullets stay.
            bottom = draw_architecture(slide, top)
            rest = [(k, p) for k, p in data.blocks if k != "para"]
            render_blocks(slide, rest, bottom + Inches(0.25), 18)
        elif text_only and len(data.blocks) == 1:
            render_hero(slide, data, top)
        else:
            render_blocks(slide, data.blocks, top, body_size(data))
        add_footer(slide, index)
    if data.notes:
        slide.notes_slide.notes_text_frame.text = "\n".join(data.notes)


def main() -> int:
    slides = parse_deck(SOURCE.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for index, data in enumerate(slides, start=1):
        render_slide(prs, data, index)
    prs.save(OUTPUT)
    print(f"wrote {OUTPUT.relative_to(HERE.parent.parent)} ({len(prs.slides)} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
